"""
Relatório DPS sem Endereço (NFSe). Base portada do MEGA QG VIVO NFSe v2.0 (legado/NFSe/qg.py);
na v3.6 saiu do porte automático (tools/portar_qgs.py) porque mudou a estrutura da planilha.
"""
from modulos.nfse.comum import *  # noqa: F401,F403 - imports do QG original

NOME_ABA1 = "DPS sem endereço"
NOME_ABA2 = "Dados de endereço"
SEPARADOR_VALORES = " | "  # quando a mesma DPS traz valores diferentes em XMLs diferentes

# (chave interna, título da coluna); as colunas de texto não podem perder zeros à esquerda
COLUNAS_ABA1 = [("dps", "DPS"), ("nome", "Nome"), ("cpf", "CPF"), ("id", "Id da DPS")]
COLUNAS_ABA2 = [("dps", "DPS"), ("nome", "Nome"), ("cpf", "CPF"), ("cep", "CEP"),
                ("cmun", "Código do município"), ("logradouro", "Logradouro"), ("numero", "Número"),
                ("complemento", "Complemento"), ("bairro", "Bairro"), ("arquivo", "Arquivo XML"),
                ("id", "Id da DPS")]
CAMPOS_TEXTO = {"dps", "cpf", "cep", "cmun", "numero", "id"}


def nome_padrao_planilha(data=None) -> str:
    """Nome padrão da planilha com a data do momento: "Planilha -08-10-2026.xlsx"."""
    return f"Planilha -{(data or datetime.now()):%d-%m-%Y}.xlsx"


def eh_nome_padrao(nome_arquivo) -> bool:
    """Se o nome é o padrão gerado (de qualquer data): aí ele é atualizado para a data de hoje."""
    import re
    return bool(re.fullmatch(r"Planilha -\d{2}-\d{2}-\d{4}\.xlsx", os.path.basename(nome_arquivo or "")))


class RelatorioSemCepEngine:
    """
    Lê os XMLs de uma pasta (tipicamente a saída "SEM CEP" do Download de XMLs) e gera uma
    planilha .xlsx com UMA LINHA POR DPS, em duas abas:

      1. "DPS sem endereço": DPS, Nome e CPF do tomador (e o Id da DPS);
      2. "Dados de endereço": a mesma DPS com os dados de endereço do tomador que
         vieram no XML (CEP, código do município, logradouro, número, complemento, bairro) e o
         nome do arquivo, para o tratamento do endereço.

    Como no QG original, só entram os XMLs com CPF e Nome no tomador. Antes as DPS eram
    agrupadas por cliente (CPF + Nome) e várias DPS iam para a mesma linha ("10, 11"). Agora
    a chave da linha é a própria DPS (o Id de <infDPS>, ou o nDPS quando não há Id): XMLs
    repetidos da mesma DPS viram uma linha só, com valores diferentes juntados na mesma célula.

    O script original usava automação COM do Excel (Excel.Application), que
    só funciona com o Excel instalado e é bem mais lento. Aqui a planilha é
    gerada com openpyxl (mesmo resultado -- duas abas, cabeçalho em negrito,
    colunas ajustadas -- só que sem precisar do Excel aberto por trás).
    Requer 'pip install openpyxl'.
    """

    def __init__(self, event_queue, stop_event, *, pasta, arquivo_saida, nome_aba1=NOME_ABA1, nome_aba2=NOME_ABA2,
                 log_dir):
        self.event_queue = event_queue
        self.stop_event = stop_event
        self.pasta = pasta
        self.arquivo_saida = arquivo_saida
        self.nome_aba1 = nome_aba1
        self.nome_aba2 = nome_aba2
        self.log_dir = log_dir
        self.log_file = None
        self.log_lock = threading.Lock()

    def emit(self, kind, **data):
        self.event_queue.put({"kind": kind, **data})

    def emit_log(self, text, status="INFO"):
        self.emit("log_line", text=text, status=status)

    def _garantir_log(self):
        os.makedirs(self.log_dir, exist_ok=True)
        self.log_file = os.path.join(self.log_dir, "relatorio_dps_sem_endereco_log_nfse.txt")

    def _log_arquivo(self, msg):
        ts = datetime.now().isoformat()
        with self.log_lock:
            with open(self.log_file, "a", encoding="utf-8") as f:
                f.write(f"[{ts}] {msg}\n")

    @staticmethod
    def _local_name(tag):
        return tag.split("}")[-1] if "}" in tag else tag

    def _find_by_localname(self, root, name):
        for elem in root.iter():
            if self._local_name(elem.tag) == name:
                return elem
        return None

    def _find_child_by_localname(self, parent, name):
        if parent is None:
            return None
        for child in parent:
            if self._local_name(child.tag) == name:
                return child
        return None

    def _texto_filho(self, parent, name):
        node = self._find_child_by_localname(parent, name)
        return node.text.strip() if node is not None and node.text else ""

    @staticmethod
    def _sanitize_sheet_name(name, fallback):
        name = (name or "").strip() or fallback
        for ch in "\\/*?:[]":
            name = name.replace(ch, "")
        return (name[:31] or fallback)

    def processar_arquivo(self, path):
        """Dados de uma DPS (dict) ou None se o tomador não tiver CPF e Nome."""
        nome_arq = os.path.basename(path)
        try:
            tree = ET.parse(path)
            root = tree.getroot()

            n_dps_node = self._find_by_localname(root, "nDPS")
            n_dps = n_dps_node.text.strip() if n_dps_node is not None and n_dps_node.text else ""
            inf_dps = self._find_by_localname(root, "infDPS")
            id_dps = (inf_dps.get("Id") or "").strip() if inf_dps is not None else ""

            toma_node = self._find_by_localname(root, "toma")
            nome = self._texto_filho(toma_node, "xNome")
            cpf = self._texto_filho(toma_node, "CPF")
            end = self._find_child_by_localname(toma_node, "end")
            end_nac = self._find_child_by_localname(end, "endNac")

            if cpf and nome:
                self._log_arquivo(f"{nome_arq} | OK | DPS={n_dps} | Nome={nome} | CPF={cpf}")
                return {
                    "id": id_dps, "dps": n_dps, "nome": nome, "cpf": cpf,
                    "cep": self._texto_filho(end_nac, "CEP"), "cmun": self._texto_filho(end_nac, "cMun"),
                    "logradouro": self._texto_filho(end, "xLgr"), "numero": self._texto_filho(end, "nro"),
                    "complemento": self._texto_filho(end, "xCpl"), "bairro": self._texto_filho(end, "xBairro"),
                    "arquivo": nome_arq,
                }
            self._log_arquivo(f"{nome_arq} | IGNORADO | sem CPF/Nome no toma")
            return None
        except Exception as e:
            self._log_arquivo(f"{nome_arq} | ERRO | {e}")
            self.emit_log(f"Erro em {nome_arq}: {e}", "ERROR")
            return None

    @staticmethod
    def consolidar(resultados):
        """Uma linha por DPS. A chave é o Id de <infDPS> (único por DPS); sem Id, o nDPS com o
        CPF. Se a mesma DPS veio em mais de um XML, os valores diferentes de cada campo ficam
        juntos na mesma célula, separados por " | ", e os arquivos são listados juntos."""
        linhas, repetidas = {}, 0
        for registro in resultados:
            chave = registro["id"] or ("nDPS", registro["dps"], registro["cpf"])
            if chave not in linhas:
                linhas[chave] = {campo: [valor] if valor else [] for campo, valor in registro.items()}
                continue
            repetidas += 1
            for campo, valor in registro.items():
                if valor and valor not in linhas[chave][campo]:
                    linhas[chave][campo].append(valor)
        juntar = {campo: (", " if campo == "arquivo" else SEPARADOR_VALORES)
                  for campo in ("id", "dps", "nome", "cpf", "cep", "cmun", "logradouro", "numero",
                                "complemento", "bairro", "arquivo")}
        consolidadas = [{campo: juntar.get(campo, SEPARADOR_VALORES).join(valores) for campo, valores in linha.items()}
                        for linha in linhas.values()]
        return consolidadas, repetidas

    def _preencher_aba(self, sheet, colunas, linhas):
        from openpyxl.styles import Font
        sheet.append([titulo for _campo, titulo in colunas])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for linha in linhas:
            sheet.append([linha.get(campo, "") for campo, _titulo in colunas])
        for indice, (campo, _titulo) in enumerate(colunas, start=1):
            if campo in CAMPOS_TEXTO:  # CPF, CEP etc. continuam texto (zeros à esquerda) se forem editados
                for cell in sheet.iter_rows(min_col=indice, max_col=indice, min_row=2):
                    cell[0].number_format = "@"
        for col_cells in sheet.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value), default=0)
            sheet.column_dimensions[col_cells[0].column_letter].width = min(length + 2, 60)
        sheet.freeze_panes = "A2"
        if linhas:
            sheet.auto_filter.ref = sheet.dimensions

    def _gerar_xlsx(self, linhas):
        try:
            import openpyxl
        except ImportError:
            raise RuntimeError(
                "A biblioteca 'openpyxl' não está instalada. Rode: pip install openpyxl")

        if os.path.exists(self.arquivo_saida):
            os.remove(self.arquivo_saida)

        wb = openpyxl.Workbook()
        sheet1 = wb.active
        sheet1.title = self._sanitize_sheet_name(self.nome_aba1, NOME_ABA1)
        self._preencher_aba(sheet1, COLUNAS_ABA1, linhas)
        sheet2 = wb.create_sheet(title=self._sanitize_sheet_name(self.nome_aba2, NOME_ABA2))
        self._preencher_aba(sheet2, COLUNAS_ABA2, linhas)

        os.makedirs(os.path.dirname(os.path.abspath(self.arquivo_saida)), exist_ok=True)
        wb.save(self.arquivo_saida)

    def run(self):
        try:
            self._garantir_log()
            arquivos = [f for f in os.listdir(self.pasta) if f.lower().endswith(".xml")]
            total = len(arquivos)
            self.emit("total_itens", n=total)
            self.emit_log(f"{total} XML(s) encontrados em {self.pasta}. Processando...", "INFO")

            resultados = []
            done = 0
            for f in arquivos:
                if self.stop_event.is_set():
                    self.emit("done", cancelled=True)
                    return
                r = self.processar_arquivo(os.path.join(self.pasta, f))
                if r:
                    resultados.append(r)
                done += 1
                self.emit("progress", done=done, total=total)

            self.emit_log(
                f"{len(resultados)} registro(s) válido(s) (com CPF e Nome) de {total} XML(s).", "INFO")

            linhas, repetidas = self.consolidar(resultados)
            self.emit_log(f"{len(linhas)} DPS na planilha (uma linha por DPS)"
                          + (f"; {repetidas} XML(s) repetido(s) da mesma DPS consolidado(s)." if repetidas else "."),
                          "INFO")

            if self.stop_event.is_set():
                self.emit("done", cancelled=True)
                return

            self.emit_log("Gerando planilha...", "INFO")
            self._gerar_xlsx(linhas)

            self.emit("done", cancelled=False, total_xml=total, validos=len(resultados),
                      agrupados=len(linhas), repetidas=repetidas, arquivo=self.arquivo_saida,
                      log_file=str(self.log_file))
        except Exception as e:
            self.emit("fatal_error", text=str(e))
