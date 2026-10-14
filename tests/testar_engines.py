"""
Testes das engines portadas e do Exportador, com HTTP e banco simulados
(não precisa de VPN, endpoint nem banco). Confere que o porte manteve o
comportamento: arquivos gerados, contadores, eventos e a URL usada por ambiente.

Rode: python tests/testar_engines.py
"""
import os
import queue
import shutil
import sys
import tempfile
import threading
from unittest import mock

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
# Preferências (tema) num arquivo temporário: os testes não mexem no tema do usuário
os.environ["MEGAQG_PREFERENCIAS"] = os.path.join(tempfile.mkdtemp(), "preferencias.json")

from core import config  # noqa: E402

falhas = []


def caso(nome):
    def decorar(fn):
        def rodar():
            pasta = tempfile.mkdtemp()
            try:
                fn(pasta)
                print(f"OK    {nome}")
            except Exception as e:  # noqa: BLE001
                falhas.append(nome)
                print(f"FALHA {nome}: {type(e).__name__}: {e}")
            finally:
                shutil.rmtree(pasta, ignore_errors=True)
        CASOS.append(rodar)
        return fn
    return decorar


CASOS = []


class Resposta:
    def __init__(self, status=200, content=b"<xml/>", text="ok", dados=None):
        self.status_code = status
        self.ok = 200 <= status < 300
        self.content = content
        self.text = text
        self._dados = dados or {}

    def json(self):
        return self._dados


def eventos(q):
    itens = []
    while not q.empty():
        itens.append(q.get())
    return itens


def config_falso(valores):
    """Substitui o config.json durante o teste."""
    return mock.patch.object(config, "_dados", valores)


@caso("NFCom Download de XMLs: baixa 3 IDs, 1 falha vai para o arquivo de erro")
def _(pasta):
    from modulos.nfcom.download_xmls.engine import DownloadXmlsEngine

    ids = os.path.join(pasta, "ids.txt")
    open(ids, "w").write("1\n2\nRUIM\n")
    erro = os.path.join(pasta, "erros.txt")

    def get(url, **_kw):
        if "RUIM" in url:
            raise ConnectionError("falhou")
        return Resposta(content=f"<xml>{url}</xml>".encode())

    q = queue.Queue()
    eng = DownloadXmlsEngine(q, threading.Event(), url_base="http://teste/download?ids={ID_DOC}",
                             pasta_saida=os.path.join(pasta, "xmls"), arquivo_erro=erro, threads=2,
                             log_dir=os.path.join(pasta, "logs"))
    with mock.patch("requests.get", side_effect=get):
        eng.run(ids)
    done = [e for e in eventos(q) if e["kind"] == "done"][0]
    assert (done["ok"], done["erro"]) == (2, 1), done
    assert sorted(os.listdir(os.path.join(pasta, "xmls"))) == ["1.xml", "2.xml"]
    assert open(erro).read().strip() == "RUIM"


@caso("NFSe Consulta DPS: URL vem do config.json conforme o ambiente")
def _(pasta):
    from modulos.nfse.consulta_dps.engine import ConsultaDpsEngine

    dps = os.path.join(pasta, "dps.txt")
    open(dps, "w").write("DPS1\nDPS2\n")
    urls = []

    def get(url, params=None, **_kw):
        urls.append(url)
        return Resposta(text="processado")

    with config_falso({"nfse_consulta_dps_url_prd": "http://prd/consulta",
                       "nfse_consulta_dps_url_hml": "http://hml/consulta"}):
        for ambiente, esperado in ((config.AMBIENTE_PRD, "http://prd/consulta"),
                                   (config.AMBIENTE_HML, "http://hml/consulta")):
            urls.clear()
            q = queue.Queue()
            eng = ConsultaDpsEngine(q, threading.Event(), ambiente=ambiente, log_dir=os.path.join(pasta, "logs"),
                                    threads=2)
            with mock.patch("requests.get", side_effect=get):
                eng.run(dps)
            assert urls and set(urls) == {esperado}, (ambiente, urls)
            done = [e for e in eventos(q) if e["kind"] == "done"][0]
            assert done["ok"] == 2, done


@caso("NFSe Download de XMLs (lote): salva os XMLs usando a URL do ambiente")
def _(pasta):
    from modulos.nfse.download_xmls.engine import DownloadXmlsEngine

    ids = os.path.join(pasta, "ids.txt")
    open(ids, "w").write("10\n11\n")
    chamadas = []

    def get(url, **kw):
        chamadas.append(url)
        return Resposta(content=b"<NFSe/>")

    with config_falso({"nfse_download_url_hml": "http://hml/download"}):
        q = queue.Queue()
        eng = DownloadXmlsEngine(q, threading.Event(), ambiente=config.AMBIENTE_HML,
                                 pasta_destino=os.path.join(pasta, "xmls"), arquivo_erro=None, threads=2,
                                 log_dir=os.path.join(pasta, "logs"))
        with mock.patch("requests.get", side_effect=get):
            eng.run(ids)
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and done["ok"] == 2, done
    assert set(chamadas) == {"http://hml/download"}, chamadas
    assert len(os.listdir(os.path.join(pasta, "xmls"))) == 2


@caso("NFSe Expandir DPS no CSV: uma linha por DPS")
def _(pasta):
    from modulos.nfse.expandir_csv.engine import ExpandirDpsCsvEngine
    import inspect

    entrada = os.path.join(pasta, "entrada.csv")
    saida = os.path.join(pasta, "saida.csv")
    open(entrada, "w", encoding="utf-8").write("Nome;DPS\nAna;1,2,3\nBeto;4\n")
    parametros = inspect.signature(ExpandirDpsCsvEngine.__init__).parameters
    kwargs = dict(arquivo_entrada=entrada, arquivo_saida=saida)
    for nome, valor in (("coluna", "DPS"), ("nome_coluna", "DPS"), ("delimitador", ";"),
                        ("separador_valores", ","), ("encoding_entrada", "utf-8"), ("encoding", "utf-8"),
                        ("log_dir", os.path.join(pasta, "logs"))):
        if nome in parametros:
            kwargs[nome] = valor
    q = queue.Queue()
    ExpandirDpsCsvEngine(q, threading.Event(), **kwargs).run()
    linhas = [linha for linha in open(saida, encoding="utf-8-sig").read().splitlines() if linha.strip()]
    assert len(linhas) == 5, linhas  # cabeçalho + 4 DPS


@caso("Exportador: período único (dia 10 a 20), só canceladas, XML em Homologação")
def _(pasta):
    import core.exportador as exportador
    import modulos.exportador_relatorio.runner as runner

    consultas, urls_xml = [], []

    def iterar(query_sql, conn, chunksize, colunas_texto):
        consultas.append(query_sql)
        yield pd.DataFrame({"ID Mensageria": [1, 2], "CNPJ Emissor": ["01", "02"]})

    def com_xml(chunk, coluna_id, url_base, max_workers, progress_callback):
        urls_xml.append(url_base)
        return chunk.assign(XML="<xml/>")

    class Conn:
        def close(self):
            pass

    with mock.patch.object(exportador, "iterar_query_em_chunks", iterar), \
         mock.patch.object(runner, "sql_conn", lambda cfg: Conn()), \
         mock.patch.object(runner, "db_config", lambda: {}), \
         mock.patch.object(runner, "adicionar_coluna_xml", com_xml), \
         config_falso({"download_url_hml": "http://hml/xml?ids={ID_DOC}"}):
        r = runner.rodar_exportacao(2026, 9, 10, 20, output_base_path=pasta, modo=runner.MODO_PERIODO,
                                    filtro_status="canceladas", ambiente_xml=config.AMBIENTE_HML)
    assert r["sucesso"] == ["10 a 20/09/2026"], r
    assert "DEMI >= '2026-09-10 00:00:00'" in consultas[0] and "DEMI < '2026-09-21 00:00:00'" in consultas[0]
    assert "IND_STATUS IN (9)" in consultas[0]
    assert urls_xml == ["http://hml/xml?ids={ID_DOC}"], urls_xml
    arquivo = os.path.join(pasta, "September", "canceladas_periodo-10-a-20-09-2026.csv")
    assert os.path.exists(arquivo), os.listdir(pasta)


@caso("Exportador: sem URL de XML configurada, avisa e não conecta no banco")
def _(pasta):
    import modulos.exportador_relatorio.runner as runner

    conectou = []
    with mock.patch.object(runner, "sql_conn", lambda cfg: conectou.append(1)), \
         config_falso({}):
        r = runner.rodar_exportacao(2026, 9, 1, 30, output_base_path=pasta, baixar_xml=True)
    assert not conectou and not r["sucesso"], r


@caso("NFCom Cancelamento: chave inválida ignorada, HTTP 500 conta como erro, XML com tpAmb e evento 110111")
def _(pasta):
    import base64
    from modulos.nfcom.cancelamento.engine import CancelamentoMassivoEngine

    boa = "35" + "2607" + "12345678000199" + "6" * 24
    ruim_http = "35" + "2607" + "12345678000199" + "7" * 24
    chaves = os.path.join(pasta, "chaves.txt")
    open(chaves, "w").write(f"{boa}\nCURTA\n{ruim_http}\n")
    enviados = []

    def post(url, headers=None, json=None, timeout=None):
        xml = base64.b64decode(json["xmlNFcom"]).decode()
        enviados.append(xml)
        return Resposta(status=500 if ruim_http in xml else 200)

    q = queue.Queue()
    eng = CancelamentoMassivoEngine(q, threading.Event(), url="http://hml/cancelar", tp_amb="2",
                                    justificativa="Teste", log_dir=os.path.join(pasta, "logs"), threads=2)
    with mock.patch("requests.post", side_effect=post):
        eng.run(chaves)
    done = [e for e in eventos(q) if e["kind"] == "done"][0]
    # Como no QG original: a chave inválida não é enviada, mas entra no contador de erro
    assert (done["ok"], done["erro"]) == (1, 2), done
    assert len(enviados) == 2
    assert all("<tpAmb>2</tpAmb>" in x and "<tpEvento>110111</tpEvento>" in x for x in enviados)


@caso("NFCom Cancelamento: sem endpoint, para antes de enviar qualquer coisa")
def _(pasta):
    from modulos.nfcom.cancelamento.engine import CancelamentoMassivoEngine

    chaves = os.path.join(pasta, "chaves.txt")
    open(chaves, "w").write("3" * 44 + "\n")
    q = queue.Queue()
    with mock.patch("requests.post") as post:
        CancelamentoMassivoEngine(q, threading.Event(), url="", tp_amb="1", justificativa="x",
                                  log_dir=os.path.join(pasta, "logs"), threads=1).run(chaves)
    assert not post.called
    assert [e for e in eventos(q) if e["kind"] == "fatal_error"]


@caso("NFCom Reenvio à SEFAZ: lê chNFCom/UF do XML e envia só os .xml da pasta")
def _(pasta):
    from modulos.nfcom.reenvio_sefaz.engine import ReenvioSefazEngine

    xmls = os.path.join(pasta, "xmls")
    os.makedirs(xmls)
    chave = "41" + "1" * 42
    open(os.path.join(xmls, "a.xml"), "w", encoding="utf-8").write(
        '<NFCom xmlns="http://www.portalfiscal.inf.br/nfcom"><infNFCom><ide/></infNFCom>'
        f'<protNFCom><infProt><chNFCom>{chave}</chNFCom></infProt></protNFCom></NFCom>')
    open(os.path.join(xmls, "leia.txt"), "w").write("nao e xml")
    corpos = []

    def post(url, data=None, headers=None, timeout=None):
        corpos.append(data)
        return Resposta(status=200, text="aceito")

    q = queue.Queue()
    eng = ReenvioSefazEngine(q, threading.Event(), url="http://hml/reenvio", threads=2,
                             log_dir=os.path.join(pasta, "logs"))
    with mock.patch("requests.post", side_effect=post):
        eng.run(xmls)
    evs = eventos(q)
    done = [e for e in evs if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and done["ok"] == 1, done
    assert len(corpos) == 1
    assert any("[PR]" in e.get("text", "") for e in evs if e["kind"] == "log_line"), "UF 41 deveria virar PR"


def _pasta_reenvio(pasta):
    """Pasta com 1 XML com chave e 5 que antes eram ignorados (o nome do arquivo é o ID no banco)."""
    xmls = os.path.join(pasta, "xmls")
    os.makedirs(xmls)
    ns = 'xmlns="http://www.portalfiscal.inf.br/nfcom"'
    arquivos = {
        "900.xml": f'<NFCom {ns}><protNFCom><infProt><chNFCom>{"35" + "1" * 42}</chNFCom></infProt></protNFCom></NFCom>',
        "101.xml": f'<NFCom {ns}><infNFCom Id="NFCom35"/></NFCom>',  # sem chNFCom -> API devolve 200
        "102.xml": f'<NFCom {ns}><infNFCom/>',                       # malformado  -> API devolve 422
        "103.xml": "<NFCom><chNFCom>35</chNFCom></NFCom>",           # sem namespace -> API devolve 500
        "nota_sem_id.xml": f'<NFCom {ns}/>',                         # nome não é ID -> continua ignorado
    }
    for nome, conteudo in arquivos.items():
        open(os.path.join(xmls, nome), "w", encoding="utf-8").write(conteudo)
    open(os.path.join(xmls, "104.xml"), "wb").write("<NFCom>São Paulo</NFCom>".encode("latin-1"))  # ilegível
    return xmls


@caso("NFCom Reenvio à SEFAZ: XML que seria ignorado é reenviado por ID (JSON [ID]) e o código HTTP vai para o log")
def _(pasta):
    from modulos.nfcom.reenvio_sefaz.engine import ReenvioSefazEngine

    xmls = _pasta_reenvio(pasta)
    codigos = {101: 200, 102: 422, 103: 500, 104: 200}
    por_id, por_xml = [], []

    def post(url, data=None, headers=None, json=None, timeout=None):
        if url == "http://hml/resend":
            por_id.append(json)
            return Resposta(status=codigos[json[0]], text='{"mensagem": "retorno da API"}')
        por_xml.append(data)
        return Resposta(status=200, text="aceito")

    q = queue.Queue()
    eng = ReenvioSefazEngine(q, threading.Event(), url="http://hml/reenvio", threads=2,
                             log_dir=os.path.join(pasta, "logs"), url_reenvio_id="http://hml/resend")
    with mock.patch("requests.post", side_effect=post):
        eng.run(xmls)
    evs = eventos(q)
    done = [e for e in evs if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done", done
    assert (done["ok"], done["erro"], done["id_ok"], done["id_erro"], done["ignorados"]) == (1, 0, 2, 2, 1), done
    assert done["total"] == 5 and done["codigos_id"] == {"HTTP 200": 2, "HTTP 422": 1, "HTTP 500": 1}, done
    assert len(por_xml) == 1 and sorted(por_id) == [[101], [102], [103], [104]], (por_xml, por_id)
    linhas = {e["text"]: e["status"] for e in evs if e["kind"] == "log_line"}
    nivel = {arquivo: next(s for t, s in linhas.items() if t.startswith("POR ID") and arquivo in t)
             for arquivo in ("101.xml", "102.xml", "103.xml")}
    assert nivel == {"101.xml": "OK", "102.xml": "WARN", "103.xml": "ERROR"}, nivel
    assert any("HTTP 422" in t and "102.xml" in t and "(ID 102" in t for t in linhas), list(linhas)
    assert any(t.startswith("IGNORADO") and "nota_sem_id.xml" in t for t in linhas), list(linhas)
    log = open(os.path.join(pasta, "logs", "integracao_sefaz_log.txt"), encoding="utf-8").read()
    assert "REENVIO_ID HTTP 422" in log and "IGNORADO" in log, log


@caso("NFCom Reenvio à SEFAZ: sem URL de reenvio por ID, os XMLs sem chave continuam ignorados")
def _(pasta):
    from modulos.nfcom.reenvio_sefaz.engine import ReenvioSefazEngine

    xmls = _pasta_reenvio(pasta)
    urls = []

    def post(url, data=None, headers=None, json=None, timeout=None):
        urls.append(url)
        return Resposta(status=200, text="aceito")

    q = queue.Queue()
    eng = ReenvioSefazEngine(q, threading.Event(), url="http://hml/reenvio", threads=2,
                             log_dir=os.path.join(pasta, "logs"))
    with mock.patch("requests.post", side_effect=post):
        eng.run(xmls)
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done", done
    assert (done["ok"], done["id_ok"], done["id_erro"], done["ignorados"]) == (1, 0, 0, 5), done
    assert urls == ["http://hml/reenvio"], urls


@caso("NFCom Reenvio à SEFAZ: reenvio por ID sem resposta da API conta como erro e segue para os próximos")
def _(pasta):
    from modulos.nfcom.reenvio_sefaz.engine import ReenvioSefazEngine

    xmls = os.path.join(pasta, "xmls")
    os.makedirs(xmls)
    for nome in ("201.xml", "202.xml"):
        open(os.path.join(xmls, nome), "w", encoding="utf-8").write("<semChave/>")

    def post(url, json=None, timeout=None):
        if json == [201]:
            raise ConnectionError("tempo esgotado")
        return Resposta(status=200)

    q = queue.Queue()
    eng = ReenvioSefazEngine(q, threading.Event(), url="http://prd/reenvio", threads=1,
                             log_dir=os.path.join(pasta, "logs"), url_reenvio_id="http://prd/resend")
    with mock.patch("requests.post", side_effect=post):
        eng.run(xmls)
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and (done["id_ok"], done["id_erro"]) == (1, 1), done
    assert done["codigos_id"] == {"HTTP 200": 1, "sem resposta": 1}, done


@caso("NFCom Reenvio de WebHook: CREATE + RESEND, conta eventos reenviados e ignorados")
def _(pasta):
    from modulos.nfcom.reenvio_webhook.engine import WebhookResendEngine

    ids = os.path.join(pasta, "ids.txt")
    open(ids, "w").write("100\n200\n")
    chamadas = []

    def post(_sessao, url, params=None, timeout=None):
        chamadas.append(url)
        if "create" in url:
            return Resposta(status=200)
        qtd = 3 if params["id-doc-fiscal"] == "100" else 0
        return Resposta(status=200, dados={"dados": {"quantidadeEventosReenviados": qtd}})

    q = queue.Queue()
    eng = WebhookResendEngine(q, threading.Event(), tipo_create="2", max_workers=2, timeout=5,
                              log_dir=os.path.join(pasta, "logs"), ambiente="Homologação")
    eng.base_url_create = "http://hml/create/{sub}"
    eng.base_url_resend = "http://hml/resend/{sub}"
    with mock.patch("requests.Session.post", post):
        eng.run(ids)
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done", done
    assert (done["ok_req"], done["ignorado_req"], done["ok_events"]) == (1, 1, 3), done
    assert "http://hml/create/1" in chamadas


@caso("NFCom Consulta SEFAZ: separa SVRS e MG, consulta todas as chaves e respeita a pausa do SVRS")
def _(pasta):
    import modulos.nfcom.consulta_sefaz.engine as mod

    chaves = ["35" + "1" * 42, "35" + "2" * 42, "35" + "3" * 42, "31" + "4" * 42]  # 3 SP (SVRS) + 1 MG
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("\n".join(chaves) + "\n")
    consultadas, esperas = [], []

    def get(url, timeout=None):
        consultadas.append(url.rsplit("/", 1)[-1])
        return Resposta(status=200, text="autorizada")

    q = queue.Queue()
    eng = mod.CheckSefazEngine(q, threading.Event(), os.path.join(pasta, "logs"), ambiente="Homologação")
    eng.base_url = "http://hml/sefaz/"
    with mock.patch("requests.get", side_effect=get), \
         mock.patch.object(mod, "time_sleep_1", lambda: esperas.append(1)):
        eng.run(arquivo)
    evs = eventos(q)
    assert [e for e in evs if e["kind"] == "all_done"], [e["kind"] for e in evs][-5:]
    total = [e for e in evs if e["kind"] == "total_keys"][0]
    assert total["grupos"]["SVRS"] == 3 and total["grupos"]["MG"] == 1, total
    assert sorted(consultadas) == sorted(chaves), consultadas
    assert len(esperas) == mod.WAIT_SVRS, esperas  # 3 chaves SVRS = 2 lotes = 1 pausa de WAIT_SVRS s


class BancoNfcomFalso:
    """documento_fiscal de mentira: {chave_nf: [id, ...]}; guarda os SELECTs feitos."""

    def __init__(self, ids):
        self.ids = ids
        self.selects = []
        self.fechado = False

    def cursor(self):
        banco = self

        class Cursor:
            def execute(self, sql, params=None):
                banco.selects.append((sql, list(params or [])))
                self.linhas = [(i, c) for c in (params or []) for i in banco.ids.get(c, [])]

            def fetchall(self):
                return self.linhas

        return Cursor()

    def close(self):
        self.fechado = True


def _consulta_com_422(pasta, chaves, respostas_get, banco, respostas_post, **kw):
    """Roda a Consulta SEFAZ com GET/POST simulados; devolve (eventos, posts)."""
    import modulos.nfcom.consulta_sefaz.engine as mod
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("\n".join(chaves) + "\n")
    posts = []

    def get(url, timeout=None):
        return Resposta(status=respostas_get.get(url.rsplit("/", 1)[-1], 200), text="retorno")

    def post(url, json=None, timeout=None):
        posts.append((url, json))
        status = respostas_post.pop(0) if respostas_post else 200
        if status is None:
            raise ConnectionError("recusada")
        return Resposta(status=status, text="processado" if status < 300 else "documento em estado inválido")

    q = queue.Queue()
    eng = mod.CheckSefazEngine(q, kw.pop("parar", threading.Event()), os.path.join(pasta, "logs"),
                               ambiente="Produção", conectar=lambda: banco)
    eng.base_url = "http://prd/sefaz/"
    eng.url_reenvio = kw.pop("url_reenvio", "http://prd/Helper/documents/resend")
    with mock.patch("requests.get", side_effect=get), mock.patch("requests.post", side_effect=post), \
         mock.patch.object(mod, "time_sleep_1", lambda: None), \
         mock.patch.object(mod, "LOTE_REENVIO", kw.pop("lote", 100)):
        eng.run(arquivo)
    return eventos(q), posts


@caso("NFCom Consulta SEFAZ: chave com HTTP 422 tem o ID buscado no banco e é reenviada com o corpo [id, id, ...]")
def _(pasta):
    c1, c2, c3, c4 = ("35" + str(n) * 42 for n in range(1, 5))  # todas SVRS
    banco = BancoNfcomFalso({c1: [101], c2: [201, 202]})  # c3 não está no banco
    evs, posts = _consulta_com_422(pasta, [c1, c2, c3, c4], {c1: 422, c2: 422, c3: 422}, banco,
                                   respostas_post=[422, 200], lote=2)
    # um SELECT só, com as 3 chaves que voltaram 422, na documento_fiscal pela chave_nf
    assert len(banco.selects) == 1 and banco.fechado, banco.selects
    sql, params = banco.selects[0]
    assert "FROM documento_fiscal" in sql and "chave_nf IN" in sql and sorted(params) == sorted([c1, c2, c3]), sql
    # IDs reenviados em lotes de 2, corpo JSON com os IDs: o 1º lote volta 422 e o reenvio continua
    ids_enviados = sorted(i for _url, corpo in posts for i in corpo)
    assert ids_enviados == [101, 201, 202] and all(url == "http://prd/Helper/documents/resend" for url, _ in posts), posts
    assert [len(corpo) for _url, corpo in posts] == [2, 1], posts
    contagem = {}
    for e in evs:
        if e["kind"] == "status_count" and e["status"].startswith("REENVIO"):
            contagem[e["status"]] = contagem.get(e["status"], 0) + 1
    assert contagem == {"REENVIO_HTTP": 2, "REENVIO_OK": 1, "REENVIO_FALHA": 1}, contagem
    linhas = [e["text"] for e in evs if e["kind"] == "log_line"]
    assert any("lote 1/2" in t and "HTTP 422" in t for t in linhas), linhas
    assert any("lote 2/2" in t and "HTTP 200" in t for t in linhas), linhas
    assert any(c3 in t and "sem ID" in t for t in linhas), linhas
    assert [e for e in evs if e["kind"] == "all_done"]
    with open(os.path.join(pasta, "logs", "checksefaz_log.txt"), encoding="utf-8") as f:
        log = f.read()
    assert f"REENVIO | {c2} | REENVIO_OK | id=202" in log and f"REENVIO | {c1} | REENVIO_HTTP | id=101" in log, log


@caso("NFCom Consulta SEFAZ: sem 422 não mexe no banco; sem URL de reenvio avisa; parada no meio não reenvia")
def _(pasta):
    c1, c2 = "35" + "1" * 42, "35" + "2" * 42
    banco = BancoNfcomFalso({c1: [101]})
    _evs, posts = _consulta_com_422(pasta, [c1, c2], {}, banco, respostas_post=[])
    assert not banco.selects and not posts  # tudo 200: nada a reenviar

    evs, posts = _consulta_com_422(pasta, [c1], {c1: 422}, banco, respostas_post=[], url_reenvio="")
    assert not banco.selects and not posts
    assert any("não configurada" in e.get("text", "") for e in evs if e["kind"] == "log_line")

    parar = threading.Event()
    parar.set()  # interrompida: as chaves não chegam a ser consultadas e nada é reenviado
    evs, posts = _consulta_com_422(pasta, [c1], {c1: 422}, banco, respostas_post=[], parar=parar)
    assert not banco.selects and not posts and [e for e in evs if e["kind"] == "all_done"]

    falhou = BancoNfcomFalso({c1: [101]})  # API fora do ar: fica no log como falha, sem derrubar a consulta
    evs, posts = _consulta_com_422(pasta, [c1], {c1: 422}, falhou, respostas_post=[None])
    assert posts == [("http://prd/Helper/documents/resend", [101])]
    assert any(e["kind"] == "status_count" and e["status"] == "REENVIO_FALHA" for e in evs)
    assert [e for e in evs if e["kind"] == "all_done"] and not [e for e in evs if e["kind"] == "fatal_error"]


@caso("NFCom banco: Produção usa as relatorio_db_* quando as nfcom_db_* estão vazias; Homologação exige as suas")
def _(pasta):
    from modulos.nfcom.banco import ConfigBancoNfcomIncompleta, db_config_nfcom
    base = {"relatorio_db_server": "srv\\PRD", "relatorio_db_database": "vivo_nfcom_prd",
            "relatorio_db_user": "u", "relatorio_db_password": "<senha>"}
    with config_falso(base):
        prd = db_config_nfcom(config.AMBIENTE_PRD)
        assert (prd["server"], prd["database"], prd["user"]) == ("srv\\PRD", "vivo_nfcom_prd", "u"), prd
        try:
            db_config_nfcom(config.AMBIENTE_HML)
            raise AssertionError("Homologação sem nfcom_db_*_hml deveria falhar")
        except ConfigBancoNfcomIncompleta as e:
            assert "nfcom_db_server_hml" in str(e) and "nfcom_db_database_hml" in str(e), e
    with config_falso({**base, "nfcom_db_server_hml": "srv\\HML", "nfcom_db_database_hml": "vivo_nfcom_hml"}):
        hml = db_config_nfcom(config.AMBIENTE_HML)
        assert (hml["server"], hml["database"], hml["password"]) == ("srv\\HML", "vivo_nfcom_hml", "<senha>"), hml


@caso("NFSe Download COM/SEM CEP: separa pelo endereço completo do tomador e do destinatário")
def _(pasta):
    from modulos.nfse.download_xmls.engine import DownloadComSemCepEngine

    end_ok = "<end><CEP>01001000</CEP><xLgr>Rua A</xLgr><nro>10</nro><xBairro>Centro</xBairro></end>"
    end_sem = "<end><CEP>00000000</CEP><xLgr>Rua A</xLgr><nro>10</nro><xBairro>Centro</xBairro></end>"
    xmls = {"1": f"<NFSe><toma>{end_ok}</toma><dest>{end_ok}</dest></NFSe>",
            "2": f"<NFSe><toma>{end_sem}</toma><dest>{end_ok}</dest></NFSe>"}
    ids = os.path.join(pasta, "ids.txt")
    open(ids, "w").write("1\n2\n")
    pedidos = []

    def get(url, params=None, **_kw):
        pedidos.append(params)
        id_nota = next(str(v) for v in (params or {}).values() if str(v) in xmls)
        return Resposta(content=xmls[id_nota].encode())

    destino = os.path.join(pasta, "saida")
    with config_falso({"nfse_download_url_prd": "http://prd/download"}):
        q = queue.Queue()
        eng = DownloadComSemCepEngine(q, threading.Event(), pasta_destino=destino, arquivo_erro=None, threads=2,
                                      log_dir=os.path.join(pasta, "logs"))
        with mock.patch("requests.get", side_effect=get):
            eng.run(ids)
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done", (done, pedidos)
    assert os.listdir(os.path.join(destino, "COM CEP")) == ["1.xml"]
    assert os.listdir(os.path.join(destino, "SEM CEP")) == ["2.xml"]


def _xml_dps_sem_endereco(id_dps, dps, nome, cpf, cep="", xlgr="", bairro="", ns=True):
    """XML de NFSe com a DPS dentro, como os que o Download de XMLs grava na pasta SEM CEP."""
    xmlns = ' xmlns="http://www.sped.fazenda.gov.br/nfse"' if ns else ""
    id_attr = f' Id="{id_dps}"' if id_dps else ""
    end = (f"<end><endNac><CEP>{cep}</CEP></endNac><xLgr>{xlgr}</xLgr><xBairro>{bairro}</xBairro></end>"
           if cep or xlgr or bairro else "")
    return (f'<NFSe{xmlns}><infNFSe><DPS><infDPS{id_attr}><nDPS>{dps}</nDPS>'
            f'<toma><CPF>{cpf}</CPF><xNome>{nome}</xNome>{end}</toma></infDPS></DPS></infNFSe></NFSe>')


@caso("NFSe Relatório DPS sem Endereço: uma linha por DPS, abas novas, endereço na 2ª aba, repetidas consolidadas")
def _(pasta):
    import openpyxl
    from modulos.nfse.relatorio_dps.engine import NOME_ABA1, NOME_ABA2, RelatorioSemCepEngine

    xmls = os.path.join(pasta, "SEM CEP")
    os.makedirs(xmls)
    notas = {
        "a.xml": _xml_dps_sem_endereco("DPS3550308200000000000101", "10", "Ana", "01234567890", "01001000",
                                       "Rua A", "Centro"),
        "b.xml": _xml_dps_sem_endereco("DPS3550308200000000000102", "11", "Ana", "01234567890"),  # mesma cliente
        "c.xml": _xml_dps_sem_endereco("DPS3550308200000000000201", "20", "Beto", "22222222222", ns=False),
        # a mesma DPS 10 de novo (outro download), com o bairro preenchido de outro jeito
        "a_copia.xml": _xml_dps_sem_endereco("DPS3550308200000000000101", "10", "Ana", "01234567890", "01001000",
                                             "Rua A", "Sé"),
        "sem_id.xml": _xml_dps_sem_endereco("", "30", "Caio", "33333333333"),
        "cnpj.xml": "<NFSe><infDPS><nDPS>40</nDPS><toma><CNPJ>1</CNPJ><xNome>Empresa</xNome></toma></infDPS></NFSe>",
    }
    for nome, conteudo in notas.items():
        open(os.path.join(xmls, nome), "w", encoding="utf-8").write(conteudo)
    saida = os.path.join(pasta, "relatorio.xlsx")
    q = queue.Queue()
    RelatorioSemCepEngine(q, threading.Event(), pasta=xmls, arquivo_saida=saida,
                          log_dir=os.path.join(pasta, "logs")).run()
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done", done
    assert (done["validos"], done["agrupados"], done["repetidas"]) == (5, 4, 1), done  # cnpj.xml fica de fora

    wb = openpyxl.load_workbook(saida)
    assert wb.sheetnames == ["DPS sem endereço", "Dados de endereço"] == [NOME_ABA1, NOME_ABA2]
    assert all(len(nome) <= 31 for nome in wb.sheetnames)
    aba1 = [list(r) for r in wb[NOME_ABA1].iter_rows(values_only=True)]
    aba2 = [list(r) for r in wb[NOME_ABA2].iter_rows(values_only=True)]
    assert aba1[0] == ["DPS", "Nome", "CPF", "Id da DPS"], aba1[0]
    assert aba2[0][:4] == ["DPS", "Nome", "CPF", "CEP"] and "Logradouro" in aba2[0] and "Bairro" in aba2[0], aba2[0]
    # 1 DPS = 1 linha, nas duas abas, sem DPS repetida e sem várias DPS na mesma célula
    assert [r[0] for r in aba1[1:]] == ["10", "11", "20", "30"], aba1
    assert [r[0] for r in aba2[1:]] == ["10", "11", "20", "30"], aba2
    assert len({r[3] for r in aba1[1:] if r[3]}) == 3  # Ids distintos (a DPS 30 não tem Id)
    ana = dict(zip(aba2[0], aba2[1]))
    assert ana["CPF"] == "01234567890" and ana["CEP"] == "01001000" and ana["Logradouro"] == "Rua A", ana
    assert ana["Bairro"] == "Centro | Sé" and ana["Arquivo XML"] == "a.xml, a_copia.xml", ana  # consolidada
    assert wb[NOME_ABA1]["C2"].number_format == "@"  # CPF continua texto (zero à esquerda)


@caso("NFSe Relatório DPS sem Endereço: nome padrão 'Planilha -dd-mm-aaaa' com a data do momento")
def _(pasta):
    from datetime import datetime as dt
    from modulos.nfse.relatorio_dps.engine import eh_nome_padrao, nome_padrao_planilha

    assert nome_padrao_planilha(dt(2026, 10, 8)) == "Planilha -08-10-2026.xlsx"
    assert nome_padrao_planilha() == f"Planilha -{dt.now():%d-%m-%Y}.xlsx"  # data real, nada fixo
    assert eh_nome_padrao(r"C:\saida\Planilha -01-01-2025.xlsx") and not eh_nome_padrao(r"C:\saida\minha.xlsx")


@caso("NFSe Reenvio de XML: envia os .xml e conta os que falharam")
def _(pasta):
    from modulos.nfse.reenvio_xml.engine import ReenvioXmlEngine

    xmls = os.path.join(pasta, "xmls")
    os.makedirs(xmls)
    for nome in ("ok.xml", "falha.xml"):
        open(os.path.join(xmls, nome), "w", encoding="utf-8").write(f"<DPS>{nome}</DPS>")

    def post(url, data=None, headers=None, timeout=None):
        return Resposta(status=500 if "falha" in data else 200)

    with config_falso({"nfse_reenvio_url_prd": "http://prd/reenvio"}):
        q = queue.Queue()
        eng = ReenvioXmlEngine(q, threading.Event(), log_dir=os.path.join(pasta, "logs"), threads=2)
        with mock.patch("requests.post", side_effect=post):
            eng.run(xmls)
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and done["ok"] == 1, done


class BancoNfceFalso:
    """Conexão pyodbc falsa: responde à prévia com situações fixas e conta commits/rollbacks."""

    def __init__(self, situacoes, falhar_bloco=None):
        self.situacoes = situacoes  # A_ID -> situação
        self.falhar_bloco = falhar_bloco
        self.commits = self.rollbacks = self.blocos = 0
        self.inserts = []

    def cursor(self):
        banco = self

        class Cursor:
            rowcount = 0
            fast_executemany = False

            def __init__(self):
                self.temp = []

            def execute(self, sql, params=None):
                if "CREATE TABLE" in sql:
                    self.temp = []
                elif "INSERT INTO COMANDO_NFEPACK" in sql:
                    banco.blocos += 1
                    banco.inserts.append((list(self.temp), list(params)))
                    if banco.falhar_bloco == banco.blocos:
                        raise RuntimeError("deadlock simulado")
                    self.rowcount = len(self.temp)

            def executemany(self, sql, linhas):
                self.temp += [linha[0] for linha in linhas]

            def fetchall(self):
                return [(c, 6, banco.situacoes.get(c, "NAO_ENCONTRADA")) for c in self.temp]

        return Cursor()

    def commit(self):
        self.commits += 1

    def rollback(self):
        self.rollbacks += 1

    def close(self):
        pass


def _engine_nfce(pasta, tipo, banco, **kw):
    from modulos.nfce.retorno_contingencia.engine import RetornoContingenciaEngine
    q = queue.Queue()
    eng = RetornoContingenciaEngine(q, threading.Event(), tipo=tipo, log_dir=os.path.join(pasta, "logs"),
                                    conectar=lambda: banco, **kw)
    return eng, q


@caso("NFCe ler_chaves: 44 dígitos válidos, inválidas e repetidas separadas")
def _(pasta):
    from modulos.nfce.retorno_contingencia.engine import ler_chaves

    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w", encoding="utf-8-sig").write("1" * 44 + "\n\n " + "2" * 44 + " \nABC\n" + "1" * 44 + "\n")
    r = ler_chaves(arquivo)
    assert r["validas"] == ["1" * 44, "2" * 44] and r["invalidas"] == ["ABC"] and r["duplicadas"] == 1, r


@caso("NFCe prévia + gravação: só as 'será inserida' vão para o INSERT, com DSC_PARAM e COMMIT")
def _(pasta):
    from modulos.nfce.retorno_contingencia.engine import IND_ACAO, STATUS_COMAND_PENDENTE

    a, b, c = "1" * 44, "2" * 44, "3" * 44
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write(f"{a}\n{b}\n{c}\n")
    banco = BancoNfceFalso({a: "SERA_INSERIDA", b: "STATUS_DIVERGENTE", c: "JA_PENDENTE"})
    eng, q = _engine_nfce(pasta, "rejeitada", banco)
    eng.previa(arquivo)
    previa = [e for e in eventos(q) if e["kind"] == "previa"][0]
    assert (previa["SERA_INSERIDA"], previa["STATUS_DIVERGENTE"], previa["JA_PENDENTE"]) == (1, 1, 1), previa
    eng.gravar(simular=False, pasta_saida=pasta)
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and done["inseridas"] == 1 and not done["simular"], done
    chaves, params = banco.inserts[0]
    assert chaves == [a], chaves
    assert params[:3] == [IND_ACAO, STATUS_COMAND_PENDENTE, "|RetornoContingencia_Rejeitada|nfepackadm"], params
    assert params[4] == 8, params  # trava de status: rejeitada = 8
    assert banco.commits == 1
    resultado = done["pasta"]
    assert open(os.path.join(resultado, "inseridas.txt")).read().split() == [a]
    assert open(os.path.join(resultado, "ja_pendentes.txt")).read().split() == [c]
    assert "GRAVAÇÃO" in open(os.path.join(resultado, "resumo.txt"), encoding="utf-8-sig").read()


@caso("NFCe simulação: executa o INSERT e desfaz (ROLLBACK), nenhum COMMIT")
def _(pasta):
    chaves = [str(i).zfill(44) for i in range(2500)]  # 3 blocos de até 1.000
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("\n".join(chaves))
    banco = BancoNfceFalso({c: "SERA_INSERIDA" for c in chaves})
    eng, q = _engine_nfce(pasta, "autorizada", banco)
    eng.previa(arquivo)
    eng.gravar(simular=True, pasta_saida=pasta)
    done = [e for e in eventos(q) if e["kind"] == "done"][0]
    assert done["simular"] and done["inseridas"] == 2500, done
    assert banco.blocos == 3 and banco.commits == 0, (banco.blocos, banco.commits)
    assert os.path.exists(os.path.join(done["pasta"], "seriam_inseridas.txt"))


@caso("NFCe: bloco com erro é desfeito e os outros blocos continuam gravados")
def _(pasta):
    chaves = [str(i).zfill(44) for i in range(2500)]
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("\n".join(chaves))
    banco = BancoNfceFalso({c: "SERA_INSERIDA" for c in chaves}, falhar_bloco=2)
    eng, q = _engine_nfce(pasta, "autorizada", banco)
    eng.previa(arquivo)
    eng.gravar(simular=False, pasta_saida=pasta)
    done = [e for e in eventos(q) if e["kind"] == "done"][0]
    assert done["inseridas"] == 1500 and done["erro"] == 1000, done
    assert banco.commits == 2, banco.commits
    assert len(open(os.path.join(done["pasta"], "erro.txt")).read().split()) == 1000


@caso("NFCe Parar: termina o bloco atual e não executa os próximos")
def _(pasta):
    chaves = [str(i).zfill(44) for i in range(2500)]
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("\n".join(chaves))
    banco = BancoNfceFalso({c: "SERA_INSERIDA" for c in chaves})
    eng, q = _engine_nfce(pasta, "autorizada", banco)
    eng.previa(arquivo)
    commit_original = banco.commit

    def commit_e_parar():
        commit_original()
        eng.stop_event.set()

    banco.commit = commit_e_parar
    eng.gravar(simular=False, pasta_saida=pasta)
    done = [e for e in eventos(q) if e["kind"] == "done"][0]
    assert done["cancelled"] and banco.blocos == 1 and done["inseridas"] == 1000, (done, banco.blocos)


@caso("NFCe: sem banco configurado no ambiente, a prévia avisa e não conecta")
def _(pasta):
    from modulos.nfce.retorno_contingencia.engine import RetornoContingenciaEngine

    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("1" * 44 + "\n")
    with config_falso({}):
        q = queue.Queue()
        RetornoContingenciaEngine(q, threading.Event(), tipo="autorizada", ambiente="Homologação",
                                  log_dir=os.path.join(pasta, "logs")).previa(arquivo)
    erro = [e for e in eventos(q) if e["kind"] == "fatal_error"]
    assert erro and "nfce_db_server_hml" in erro[0]["text"], erro


@caso("NFCe Download de XMLs: salva <chave>.xml, HTTP 404/500 e resposta sem XML vão para o arquivo de erro")
def _(pasta):
    from modulos.nfce.download_xmls.engine import DownloadXmlsNfceEngine

    boa, em_json, nao_achou, quebrou, sem_xml = ("1" * 44, "2" * 44, "3" * 44, "4" * 44, "5" * 44)
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("\n".join([boa, em_json, nao_achou, quebrou, sem_xml, "CURTA", boa]) + "\n")
    respostas = {
        boa: Resposta(content=b"<nfeProc>boa</nfeProc>"),
        em_json: Resposta(content=b'"<nfeProc>em json</nfeProc>"'),  # API que devolve o XML como string JSON
        nao_achou: Resposta(status=404, content=b"", text="nao encontrada"),
        quebrou: Resposta(status=500, content=b"<html>erro</html>", text="erro interno"),
        sem_xml: Resposta(content=b"", text=""),
    }
    urls = []

    def get(url, timeout=None, verify=None):
        urls.append(url)
        return respostas[url.rsplit("=", 1)[-1]]

    saida = os.path.join(pasta, "xmls")
    with config_falso({"nfce_download_url_prd": "http://prd/nfce/api/XML/XMLEmissao?chave="}):
        q = queue.Queue()
        eng = DownloadXmlsNfceEngine(q, threading.Event(), pasta_saida=saida, threads=2,
                                     log_dir=os.path.join(pasta, "logs"))
        with mock.patch("requests.get", side_effect=get):
            eng.run(arquivo)
    evs = eventos(q)
    done = [e for e in evs if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and (done["ok"], done["erro"], done["invalidas"]) == (2, 3, 1), done
    assert sorted(urls) == sorted(f"http://prd/nfce/api/XML/XMLEmissao?chave={c}" for c in respostas), urls
    assert open(os.path.join(saida, f"{boa}.xml"), "rb").read() == b"<nfeProc>boa</nfeProc>"
    assert open(os.path.join(saida, f"{em_json}.xml"), "rb").read() == b"<nfeProc>em json</nfeProc>"
    assert sorted(f for f in os.listdir(saida) if f.endswith(".xml")) == [f"{boa}.xml", f"{em_json}.xml"]
    assert sorted(open(done["arquivo_erro"]).read().split()) == [nao_achou, quebrou, sem_xml]
    textos = [e["text"] for e in evs if e["kind"] == "log_line"]
    assert any("HTTP 404" in t and nao_achou in t for t in textos), textos
    assert any("HTTP 500" in t and quebrou in t for t in textos), textos


@caso("NFCe Download de XMLs: URL do ambiente com {CHAVE}; sem URL, avisa e não chama a API")
def _(pasta):
    from modulos.nfce.download_xmls.engine import DownloadXmlsNfceEngine

    chave = "7" * 44
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write(chave + "\n")
    urls = []

    def get(url, timeout=None, verify=None):
        urls.append(url)
        return Resposta(content=b"<nfeProc/>")

    with config_falso({"nfce_download_url_hml": "http://hml/xml/{CHAVE}/emissao"}), \
         mock.patch("requests.get", side_effect=get):
        for ambiente in (config.AMBIENTE_HML, config.AMBIENTE_PRD):
            q = queue.Queue()
            DownloadXmlsNfceEngine(q, threading.Event(), ambiente=ambiente, pasta_saida=os.path.join(pasta, "xmls"),
                                   threads=1, log_dir=os.path.join(pasta, "logs")).run(arquivo)
            fim = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
            if ambiente == config.AMBIENTE_HML:
                assert fim["kind"] == "done" and fim["ok"] == 1, fim
            else:
                assert fim["kind"] == "fatal_error" and "nfce_download_url_prd" in fim["text"], fim
    assert urls == [f"http://hml/xml/{chave}/emissao"], urls


@caso("Registro: NFCe tem o Retorno de Contingência numa tela só e o Download de XMLs, que não exige o banco")
def _(pasta):
    from modulos.registro import OPERACOES, carregar_pagina, pendencias

    nfce = {op["id"]: op for op in OPERACOES if op["area"] == "NFCe"}
    assert list(nfce) == ["nfce.retorno_contingencia", "nfce.download_xmls", "nfce.reprocessamento"], list(nfce)
    for op in nfce.values():
        carregar_pagina(op)  # a classe da tela existe
    with config_falso({"nfce_download_url_prd": "http://prd/xml?chave="}):
        assert pendencias(nfce["nfce.download_xmls"], config.AMBIENTE_PRD) == []
        assert pendencias(nfce["nfce.download_xmls"], config.AMBIENTE_HML) == ["nfce_download_url_hml"]
        assert "nfce_db_user" in pendencias(nfce["nfce.retorno_contingencia"], config.AMBIENTE_PRD)
    reenvio = next(op for op in OPERACOES if op["id"] == "nfcom.reenvio_sefaz")
    with config_falso({"reenvio_sefaz_url_prd": "http://prd/reenvio"}):
        assert pendencias(reenvio, config.AMBIENTE_PRD) == ["reenvio_sefaz_id_url_prd"]


class BancoReprocessamentoFalso:
    """Conexão pyodbc falsa para o Reprocessamento: INTERF_NFE em memória (A_ID -> [NUM_SEQ_NFE, status]).
    O UPDATE só vale depois do COMMIT; o ROLLBACK descarta."""

    def __init__(self, notas, falhar_update=False):
        self.notas = {chave: list(valor) for chave, valor in notas.items()}
        self.falhar_update = falhar_update
        self.commits = self.rollbacks = 0
        self.updates = []  # NUM_SEQ_NFE de cada UPDATE executado
        self._pendente = []

    def cursor(self):
        banco = self

        class Cursor:
            rowcount = 0
            fast_executemany = False

            def __init__(self):
                self.temp, self.seq, self.resultado = [], [], []

            def execute(self, sql, params=None):
                if "CREATE TABLE" in sql:
                    self.temp = []
                elif "INTO #seq" in sql:
                    pular = "(1 = 0" in sql
                    self.seq = [(c, banco.notas[c][0]) for c in self.temp
                                if c in banco.notas and not (pular and banco.notas[c][1] == 6)]
                elif sql.startswith("UPDATE INTERF_NFE"):
                    if banco.falhar_update:
                        raise RuntimeError("deadlock simulado")
                    assert "IND_STATUS_NFE = 3" in sql and "NUM_SEQ_NFE IN" in sql, sql
                    banco.updates.append([seq for _chave, seq in self.seq])
                    banco._pendente = [seq for _chave, seq in self.seq]
                    self.rowcount = len(self.seq)
                elif "FROM #seq" in sql:
                    self.resultado = list(self.seq)
                elif "FROM #chaves" in sql:
                    self.resultado = [(c, *banco.notas.get(c, (None, None))) for c in self.temp]

            def executemany(self, sql, linhas):
                self.temp += [linha[0] for linha in linhas]

            def fetchall(self):
                return self.resultado

        return Cursor()

    def commit(self):
        self.commits += 1
        for nota in self.notas.values():
            if nota[0] in self._pendente:
                nota[1] = 3
        self._pendente = []

    def rollback(self):
        self.rollbacks += 1
        self._pendente = []

    def close(self):
        pass


def _xml_nfce(chave):
    """XML como o download devolve: o <NFe> assinado dentro de um nfeProc."""
    ns = 'xmlns="http://www.portalfiscal.inf.br/nfe"'
    return (f'<?xml version="1.0" encoding="UTF-8"?><nfeProc {ns} versao="4.00"><NFe {ns}>'
            f'<infNFe versao="4.00" Id="NFe{chave}"><ide><mod>65</mod></ide></infNFe><Signature/></NFe>'
            f'<protNFe><infProt><chNFe>{chave}</chNFe></infProt></protNFe></nfeProc>').encode()


CONFIG_REPROCESSAMENTO = {"nfce_download_url_prd": "http://prd/nfce/api/XML/XMLEmissao?chave=",
                          "nfce_reprocessamento_url_prd": "http://prd/nfce/api/nfe/RECEPCIONARNOTASINCRONA"}


def _reprocessamento(pasta, notas, respostas_envio=None, sem_xml=(), **kw):
    """Roda prévia + reprocessamento com banco e HTTP simulados. Devolve (engine, banco, eventos, envios)."""
    from modulos.nfce.reprocessamento.engine import ReprocessamentoNfceEngine

    simular = kw.pop("simular", False)
    falhar_update = kw.pop("falhar_update", False)
    parar_no_envio = kw.pop("parar_no_envio", False)  # simula o botão Parar depois do 1º envio
    parar = threading.Event()
    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("\n".join(kw.pop("chaves")) + "\n")
    banco = BancoReprocessamentoFalso(notas, falhar_update=falhar_update)
    envios = {}

    def get(url, timeout=None, verify=None):
        chave = url.rsplit("=", 1)[-1]
        if chave in sem_xml:
            return Resposta(status=404, content=b"", text="nao encontrada")
        return Resposta(content=_xml_nfce(chave))

    def post(url, data=None, headers=None, timeout=None, verify=None):
        assert url == CONFIG_REPROCESSAMENTO["nfce_reprocessamento_url_prd"], url
        assert headers == {"Content-Type": "application/xml"}, headers
        chave = data.split(b'Id="NFe')[1][:44].decode()
        envios[chave] = data
        if parar_no_envio:
            parar.set()
        return (respostas_envio or {}).get(chave, Resposta(status=200, text="<cStat>100</cStat>"))

    q = queue.Queue()
    with config_falso(CONFIG_REPROCESSAMENTO), mock.patch("requests.get", side_effect=get), \
         mock.patch("requests.post", side_effect=post):
        eng = ReprocessamentoNfceEngine(q, parar, log_dir=os.path.join(pasta, "logs"),
                                        conectar=lambda: banco, threads=kw.pop("threads", 2), **kw)
        eng.previa(arquivo)
        eng.reprocessar(simular=simular, pasta_saida=pasta)
    return eng, banco, eventos(q), envios


@caso("NFCe Reprocessamento: separa o <NFe> do XML baixado sem alterar os bytes e lê cStat/xMotivo do retorno")
def _(pasta):
    from modulos.nfce.reprocessamento.engine import extrair_nfe, retorno_sefaz

    ns = 'xmlns="http://www.portalfiscal.inf.br/nfe"'
    declaracao = b'<?xml version="1.0" encoding="utf-8"?>'
    nfe = f'<NFe {ns}><infNFe Id="NFe1"><xNome>São João</xNome></infNFe><Signature/></NFe>'.encode()
    assert extrair_nfe(b'<?xml version="1.0"?><nfeProc versao="4.00">' + nfe + b"<protNFe/></nfeProc>") == declaracao + nfe
    assert extrair_nfe(nfe) == declaracao + nfe  # já veio só o <NFe>
    assert extrair_nfe(f"<nfeProc {ns}><NFe><infNFe/></NFe></nfeProc>".encode()) == \
        declaracao + f"<NFe {ns}><infNFe/></NFe>".encode()  # xmlns só no elemento de fora
    assert extrair_nfe(b"<erro>nota nao encontrada</erro>") is None and extrair_nfe(None) is None
    assert retorno_sefaz("<retEnviNFe><cStat>104</cStat><xMotivo>Lote processado</xMotivo><protNFe><infProt>"
                         "<cStat>100</cStat><xMotivo>Autorizado o uso da NF-e</xMotivo></infProt></protNFe>"
                         "</retEnviNFe>") == "cStat 100 Autorizado o uso da NF-e"
    assert retorno_sefaz('{"cStat": "539", "xMotivo": "Rejeicao: Duplicidade"}') == "cStat 539 Rejeicao: Duplicidade"
    assert retorno_sefaz("Internal Server Error") == ""


@caso("NFCe Reprocessamento: baixa o XML, grava status 3 só para quem tem XML e envia o <NFe> completo")
def _(pasta):
    a, b, c, d, e = ("1" * 44, "2" * 44, "3" * 44, "4" * 44, "5" * 44)
    notas = {a: (11, 8), b: (12, 4), c: (13, 8), d: (14, 6)}  # d já autorizada; e não existe no banco
    eng, banco, evs, envios = _reprocessamento(
        pasta, notas, chaves=[a, b, c, d, e], sem_xml={c},
        respostas_envio={a: Resposta(status=200, text="<cStat>100</cStat><xMotivo>Autorizado o uso da NF-e</xMotivo>"),
                         b: Resposta(status=500, text="Internal Server Error")})
    previa = [x for x in evs if x["kind"] == "previa"][0]
    assert (previa["SERA_REPROCESSADA"], previa["JA_AUTORIZADA"], previa["NAO_ENCONTRADA"]) == (3, 1, 1), previa
    assert previa["status_atual"] == {"4": 1, "6": 1, "8": 2}, previa
    done = [x for x in evs if x["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done", done
    assert (done["baixados"], done["atualizadas"], done["enviadas"], done["erro"]) == (2, 2, 1, 2), done
    assert done["codigos"] == {"HTTP 200": 1, "HTTP 500": 1}, done
    # o UPDATE usa o NUM_SEQ_NFE das notas com XML; a que falhou no download não muda de status
    assert banco.updates == [[11, 12]] and banco.commits == 1, (banco.updates, banco.commits)
    assert {k: v[1] for k, v in banco.notas.items()} == {a: 3, b: 3, c: 8, d: 6}, banco.notas
    # corpo do POST: só o <NFe>, com a declaração XML, sem o nfeProc de fora
    assert sorted(envios) == [a, b], sorted(envios)
    assert envios[a].startswith(b'<?xml version="1.0" encoding="utf-8"?><NFe xmlns='), envios[a][:80]
    assert envios[a].endswith(b"</NFe>") and b"nfeProc" not in envios[a] and b"protNFe" not in envios[a]
    resultado = done["pasta"]
    lista = lambda nome: open(os.path.join(resultado, nome)).read().split()  # noqa: E731
    assert lista("enviadas.txt") == [a] and lista("erro_envio.txt") == [b] and lista("erro_download.txt") == [c]
    assert lista("ja_autorizadas.txt") == [d] and lista("nao_encontradas.txt") == [e]
    retornos = open(os.path.join(resultado, "retornos.txt"), encoding="utf-8-sig").read()
    assert f"{a} | 11 | HTTP 200 | cStat 100 Autorizado o uso da NF-e" in retornos, retornos
    assert f"{b} | 12 | HTTP 500" in retornos, retornos
    assert open(os.path.join(resultado, "xmls", f"{a}.xml"), "rb").read() == envios[a]
    niveis = {x["text"].split()[2]: x["status"] for x in evs if x["kind"] == "log_line" and x["text"].startswith("HTTP")}
    assert niveis == {a: "OK", b: "ERROR"}, niveis


@caso("NFCe Reprocessamento, simulação: baixa os XMLs, UPDATE com ROLLBACK e nenhum envio")
def _(pasta):
    chaves = [str(i).zfill(44) for i in range(450)]  # 3 blocos de até 200
    notas = {chave: (1000 + i, 8) for i, chave in enumerate(chaves)}
    eng, banco, evs, envios = _reprocessamento(pasta, notas, chaves=chaves, simular=True)
    done = [x for x in evs if x["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and done["simular"], done
    assert (done["baixados"], done["atualizadas"], done["enviadas"], done["erro"]) == (450, 450, 0, 0), done
    assert len(banco.updates) == 3 and banco.commits == 0 and not envios, (len(banco.updates), banco.commits)
    assert all(status == 8 for _seq, status in banco.notas.values())
    assert len(open(os.path.join(done["pasta"], "seriam_reprocessadas.txt")).read().split()) == 450


@caso("NFCe Reprocessamento: UPDATE com erro é desfeito e nada do bloco é enviado; nota autorizada só entra sem a trava")
def _(pasta):
    a, d = "1" * 44, "4" * 44
    eng, banco, evs, envios = _reprocessamento(os.path.join(pasta), {a: (11, 8)}, chaves=[a], falhar_update=True)
    done = [x for x in evs if x["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and (done["enviadas"], done["erro"]) == (0, 1), done
    assert not envios and banco.commits == 0 and banco.notas[a][1] == 8
    assert open(os.path.join(done["pasta"], "erro_atualizacao.txt")).read().split() == [a]

    eng, banco, evs, envios = _reprocessamento(pasta, {d: (14, 6)}, chaves=[d], pular_autorizadas=False)
    done = [x for x in evs if x["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and done["enviadas"] == 1 and banco.notas[d][1] == 3, (done, banco.notas)


@caso("NFCe Reprocessamento, Parar no envio: lista as notas que ficaram com status 3 sem envio e não abre o bloco seguinte")
def _(pasta):
    chaves = [str(i).zfill(44) for i in range(250)]  # bloco 1 = 200, bloco 2 = 50
    notas = {chave: (1000 + i, 8) for i, chave in enumerate(chaves)}
    eng, banco, evs, envios = _reprocessamento(pasta, notas, chaves=chaves, parar_no_envio=True, threads=1)
    done = [x for x in evs if x["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and done["cancelled"], done
    assert (len(envios), done["enviadas"], done["pendentes"]) == (1, 1, 199), (len(envios), done)
    assert len(banco.updates) == 1 and banco.commits == 1, (len(banco.updates), banco.commits)  # bloco 2 intocado
    lista = lambda nome: open(os.path.join(done["pasta"], nome)).read().split()  # noqa: E731
    assert len(lista("status3_nao_enviadas.txt")) == 199 and lista("nao_processadas.txt") == chaves[200:]
    assert sum(1 for _seq, status in banco.notas.values() if status == 3) == 200


@caso("NFCe Reprocessamento: sem as URLs do ambiente, a prévia avisa e não conecta no banco")
def _(pasta):
    from modulos.nfce.reprocessamento.engine import ReprocessamentoNfceEngine

    arquivo = os.path.join(pasta, "chaves.txt")
    open(arquivo, "w").write("1" * 44 + "\n")
    conectou = []
    with config_falso({"nfce_download_url_prd": "http://prd/xml?chave="}):
        q = queue.Queue()
        ReprocessamentoNfceEngine(q, threading.Event(), log_dir=os.path.join(pasta, "logs"),
                                  conectar=lambda: conectou.append(1)).previa(arquivo)
    erro = [x for x in eventos(q) if x["kind"] == "fatal_error"]
    assert erro and "nfce_reprocessamento_url_prd" in erro[0]["text"] and not conectou, erro


NS_NFSE = "http://www.sped.fazenda.gov.br/nfse"


def _estragar(texto, vezes=1, como="cp1252"):
    """Simula o erro de conversão: UTF-8 lido como Windows-1252 (ou Latin-1), `vezes` camadas."""
    for _ in range(vezes):
        texto = texto.encode("utf-8").decode(como)
    return texto


def _dps(corpo, xmlns=True):
    """DPS mínima com o conteúdo de <infDPS> informado (com xmlns, como os XMLs reais)."""
    import xml.etree.ElementTree as ET
    atributo = f' xmlns="{NS_NFSE}"' if xmlns else ""
    return ET.fromstring(f'<DPS{atributo} versao="1.00"><infDPS Id="DPS1">{corpo}</infDPS></DPS>')


def _tag(raiz, nome):
    from modulos.nfse.correcao_xml.correcoes import descendente
    return descendente(raiz, nome)


def _nomes(elemento):
    from modulos.nfse.correcao_xml.correcoes import nome_local
    return [nome_local(filho.tag) for filho in elemento]


def _apis(respostas):
    """requests.get falso: escolhe a resposta pelo trecho da URL; o que não estiver listado dá 404."""
    chamadas = []

    def get(url, **_kw):
        chamadas.append(url)
        for trecho, resposta in respostas.items():
            if trecho in url:
                if isinstance(resposta, Exception):
                    raise resposta
                return resposta
        return Resposta(status=404)

    return get, chamadas


@caso("NFSe Correção de XMLs: caracteres corrompidos só nas tags x*, em mais de uma camada, sem estragar texto bom")
def _(pasta):
    from modulos.nfse.correcao_xml.correcoes import corrigir_encoding, corrigir_texto_corrompido

    assert corrigir_texto_corrompido(_estragar("não", 2)) == ("não", True)
    assert corrigir_texto_corrompido(_estragar("Anúncios", 2)) == ("Anúncios", True)
    assert corrigir_texto_corrompido(_estragar("Érico")) == ("Érico", True)
    assert corrigir_texto_corrompido(_estragar("Índio", como="latin-1")) == ("Índio", True)
    assert corrigir_texto_corrompido("Condominio " + _estragar("NÃO") + " INFORMADO") == (
        "Condominio NÃO INFORMADO", True)
    assert corrigir_texto_corrompido(_estragar("UNIÃO")) == ("UNIÃO", True)
    for bom in ("SÃO JOÃO DA BARRA", "Avenida Brasil, 100", "Ação e Reação Ltda", ""):
        assert corrigir_texto_corrompido(bom) == (bom, False), bom

    raiz = _dps(f"<toma><xNome>{_estragar('José Conceição', 2)}</xNome><CNPJ>{_estragar('ã')}</CNPJ>"
                f"<end><xLgr>Rua {_estragar('São Bento')}</xLgr><xBairro>Centro</xBairro></end></toma>")
    quantidade, mensagens = corrigir_encoding(raiz)
    assert quantidade == 2 and len(mensagens) == 3, mensagens
    assert _tag(raiz, "xNome").text == "José Conceição" and _tag(raiz, "xLgr").text == "Rua São Bento"
    assert _tag(raiz, "CNPJ").text == _estragar("ã")  # tag que não começa com "x" não é tocada


@caso("NFSe Correção de XMLs: <nro> criado só quando a tag não existe; competência segue a data de <dhEmi>")
def _(pasta):
    from modulos.nfse.correcao_xml.correcoes import corrigir_competencia, preencher_numero_ausente

    raiz = _dps("<dhEmi>2026-07-20T08:30:00-03:00</dhEmi><dCompet>2026-07-19</dCompet>"
                "<toma><end><xLgr>Rua A</xLgr><xBairro>Centro</xBairro></end></toma>"
                "<dest><end><xCpl>Sala 2</xCpl><xBairro>Centro</xBairro></end></dest>"
                "<interm><end><xLgr>Rua C</xLgr></end></interm>")
    assert preencher_numero_ausente(raiz)[0] == 2
    assert _nomes(_tag(_tag(raiz, "toma"), "end")) == ["xLgr", "nro", "xBairro"]  # logo após <xLgr>
    assert _nomes(_tag(_tag(raiz, "dest"), "end")) == ["nro", "xCpl", "xBairro"]  # sem xLgr: antes de xCpl
    assert _nomes(_tag(_tag(raiz, "interm"), "end")) == ["xLgr"]  # só tomador e destinatário
    assert _tag(raiz, "nro").text == "01" and NS_NFSE in _tag(raiz, "nro").tag  # herda o xmlns
    assert preencher_numero_ausente(raiz)[0] == 0
    vazio = _dps("<toma><end><xLgr>Rua A</xLgr><nro></nro></end></toma>")
    assert preencher_numero_ausente(vazio)[0] == 0 and not _tag(vazio, "nro").text  # <nro> vazio não é alterado

    assert corrigir_competencia(raiz)[0] == 1 and _tag(raiz, "dCompet").text == "2026-07-20"
    assert corrigir_competencia(raiz)[0] == 0  # já igual
    assert corrigir_competencia(_dps("<dhEmi>2026-07-20T08:30:00-03:00</dhEmi>"))[0] == 0  # não cria <dCompet>
    invalida = _dps("<dhEmi>20/07/2026</dhEmi><dCompet>2026-07-19</dCompet>")
    assert corrigir_competencia(invalida)[0] == 0 and _tag(invalida, "dCompet").text == "2026-07-19"
    virada = _dps("<dhEmi>2026-08-01T00:10:00-03:00</dhEmi><dCompet></dCompet>")  # sem conversão de fuso
    assert corrigir_competencia(virada)[0] == 1 and _tag(virada, "dCompet").text == "2026-08-01"


@caso("NFSe Correção de XMLs: PIS/COFINS (CST ausente, CST sem base, alíquota ×100, valor recalculado) e ISS")
def _(pasta):
    from modulos.nfse.correcao_xml.correcoes import corrigir_aliquota_iss, corrigir_piscofins

    sem_cst = _dps("<piscofins><vBCPisCofins>100.00</vBCPisCofins><pAliqPis>1.65</pAliqPis>"
                   "<pAliqCofins>7.60</pAliqCofins><vPis>1.65</vPis><vCofins>7.60</vCofins></piscofins>")
    assert corrigir_piscofins(sem_cst)[0] == 1
    assert _nomes(_tag(sem_cst, "piscofins"))[0] == "CST" and _tag(sem_cst, "CST").text == "01"
    assert corrigir_piscofins(_dps("<piscofins><CST></CST><vPis>1.00</vPis></piscofins>"), "02")[0] == 1

    sem_base = _dps("<piscofins><CST>8</CST><vBCPisCofins>100.00</vBCPisCofins><pAliqPis>1.65</pAliqPis>"
                    "<pAliqCofins>7.60</pAliqCofins><vPis>1.65</vPis><vCofins>7.60</vCofins></piscofins>")
    assert corrigir_piscofins(sem_base)[0] == 6  # CST padronizado + 5 campos removidos (E0682)
    assert _nomes(_tag(sem_base, "piscofins")) == ["CST"] and _tag(sem_base, "CST").text == "08"

    # alíquotas ×100 por engano e vPis que o arredondamento comum erra: 107,00 × 2,50% = 2,675 -> 2,68
    tributado = _dps("<piscofins><CST>01</CST><vBCPisCofins>107.00</vBCPisCofins><pAliqPis>250</pAliqPis>"
                     "<pAliqCofins>760</pAliqCofins><vPis>2.67</vPis><vCofins>8.13</vCofins></piscofins>")
    quantidade, mensagens = corrigir_piscofins(tributado)
    assert quantidade == 3, mensagens  # duas alíquotas e o vPis; vCofins (8,132 -> 8,13) já estava certo
    valores = {nome: _tag(tributado, nome).text for nome in ("pAliqPis", "pAliqCofins", "vPis", "vCofins")}
    assert valores == {"pAliqPis": "2.50", "pAliqCofins": "7.60", "vPis": "2.68", "vCofins": "8.13"}, valores
    assert corrigir_piscofins(tributado)[0] == 0

    try:
        corrigir_piscofins(_dps("<piscofins><xObs>sem nada</xObs></piscofins>"))
        raise AssertionError("sem CST e sem dados deveria virar pendente")
    except RuntimeError as erro:
        assert "sem CST" in str(erro)

    iss = _dps("<tribMun><pAliq>500</pAliq></tribMun><outro><pAliq>5.00</pAliq></outro>")
    assert corrigir_aliquota_iss(iss)[0] == 1
    assert [e.text for e in iss.iter() if e.tag.endswith("pAliq")] == ["5.00", "5.00"]


@caso("NFSe Correção de XMLs: cMun pelo CEP (APIs em ordem, IBGE, caches) e destinatário reaproveitando o do tomador")
def _(pasta):
    from modulos.nfse.correcao_xml.cep import ResolvedorMunicipio
    from modulos.nfse.correcao_xml.correcoes import ErroCorrecaoCep, corrigir_cmun

    def endereco(cep, cmun=None):
        cmun = "" if cmun is None else f"<cMun>{cmun}</cMun>"
        return f"<end><endNac>{cmun}<CEP>{cep}</CEP></endNac><xLgr>Rua A</xLgr></end>"

    get, chamadas = _apis({
        "viacep.com.br/ws/01001000": Resposta(dados={"ibge": "3550308", "localidade": "São Paulo", "uf": "SP"}),
        "viacep.com.br/ws/30130010": Resposta(dados={"localidade": "Belo Horizonte", "uf": "MG"}),  # sem IBGE
        "localidades/estados/MG": Resposta(dados=[{"id": 3106200, "nome": "Belo Horizonte"}]),
        "viacep.com.br/ws/99999999": Resposta(dados={"erro": "true"}),
        "88888888": ConnectionError("sem rede"),
        "viacep.com.br/ws/69900000": Resposta(status=500),
        "brasilapi.com.br/api/cep/v2/69900000": Resposta(dados={"city": "Rio Branco", "state": "AC"}),
        "localidades/estados/AC": Resposta(dados=[{"id": 1200401, "nome": "Rio Branco"}]),
    })
    with mock.patch("requests.get", side_effect=get):
        resolvedor = ResolvedorMunicipio(os.path.join(pasta, "dados"), pausa=0)
        # tomador sem cMun; destinatário com o mesmo CEP reaproveita; 2º destinatário tem CEP próprio
        raiz = _dps(f"<toma>{endereco('01001-000')}</toma><dest>{endereco('01001000')}</dest>"
                    f"<toma>{endereco('30130-010', '')}</toma><dest>{endereco('69900000')}</dest>")
        quantidade, mensagens = corrigir_cmun(raiz, resolvedor, {})
        assert quantidade == 4, mensagens
        assert [e.text for e in raiz.iter() if e.tag.endswith("cMun")] == ["3550308", "3550308", "3106200", "1200401"]
        assert _nomes(_tag(raiz, "endNac")) == ["cMun", "CEP"]  # criado imediatamente antes de <CEP>
        assert sum("01001000" in url for url in chamadas) == 1  # o destinatário não consultou de novo
        assert any("cep/v2/69900000" in url for url in chamadas)  # ViaCEP falhou: caiu para a BrasilAPI

        preenchido = _dps(f"<toma>{endereco('12345678', '4106902')}</toma>")
        antes = len(chamadas)
        assert corrigir_cmun(preenchido, resolvedor, {})[0] == 0 and len(chamadas) == antes  # não consulta

        for cep, entra_no_cache in (("99999999", True), ("88888888", False)):
            try:
                corrigir_cmun(_dps(f"<toma>{endereco(cep)}</toma>"), resolvedor, {})
                raise AssertionError("CEP sem resolução deveria virar pendente")
            except ErroCorrecaoCep:
                assert (cep in resolvedor.invalidos) == entra_no_cache, cep  # falha de rede não é "inválido"
        try:
            corrigir_cmun(_dps("<toma><end><endNac></endNac></end></toma>"), resolvedor, {})
            raise AssertionError("sem CEP deveria virar pendente")
        except ErroCorrecaoCep as erro:
            assert "não possui CEP" in str(erro)

        # caches em arquivo: outra execução não consulta de novo o CEP resolvido nem o inválido
        antes = len(chamadas)
        novo = ResolvedorMunicipio(os.path.join(pasta, "dados"), pausa=0)
        assert novo.cmun_por_cep("01001000") == "3550308" and "99999999" in novo.invalidos
        try:
            novo.cmun_por_cep("99999999")
            raise AssertionError("CEP inválido em cache deveria falhar")
        except RuntimeError as erro:
            assert "cache de CEPs inválidos" in str(erro)
        assert len(chamadas) == antes


@caso("NFSe Correção de XMLs: cLocEmi e cLocPrestacao pelo cMun do emitente, com xmlns e pelo CNPJ")
def _(pasta):
    from modulos.nfse.correcao_xml.cep import ResolvedorMunicipio
    from modulos.nfse.correcao_xml.correcoes import corrigir_clocemi, corrigir_clocprestacao

    prest = "<prest><CNPJ>12345678000199</CNPJ><end><endNac><cMun>3106200</cMun><CEP>30130010</CEP></endNac></end></prest>"
    raiz = _dps(f"<tpAmb>1</tpAmb><tpEmit>1</tpEmit>{prest}<serv><cServ><cTribNac>010101</cTribNac></cServ></serv>")
    assert corrigir_clocemi(raiz, None)[0] == 1 and corrigir_clocprestacao(raiz, None)[0] == 1
    assert _nomes(_tag(raiz, "infDPS")) == ["tpAmb", "tpEmit", "cLocEmi", "prest", "serv"]  # logo após <tpEmit>
    assert _nomes(_tag(raiz, "serv")) == ["locPrest", "cServ"]  # <locPrest> criado antes de <cServ>
    assert _tag(raiz, "cLocEmi").text == _tag(raiz, "cLocPrestacao").text == "3106200"
    assert all(NS_NFSE in _tag(raiz, nome).tag for nome in ("cLocEmi", "locPrest", "cLocPrestacao"))
    assert corrigir_clocemi(raiz, None)[0] == 0 and corrigir_clocprestacao(raiz, None)[0] == 0  # já preenchidos

    zerado = _dps(f"<tpEmit>1</tpEmit><cLocEmi>0</cLocEmi>{prest}<serv><locPrest><cLocPrestacao/></locPrest></serv>")
    assert corrigir_clocemi(zerado, None)[0] == 1 and corrigir_clocprestacao(zerado, None)[0] == 1
    assert _tag(zerado, "cLocEmi").text == _tag(zerado, "cLocPrestacao").text == "3106200"

    # emitente sem cMun no XML: CNPJ -> CEP -> cMun
    get, chamadas = _apis({"brasilapi.com.br/api/cnpj/v1/12345678000199": Resposta(dados={"cep": "30130-010"}),
                           "viacep.com.br/ws/30130010": Resposta(dados={"ibge": "3106200"})})
    pelo_cnpj = _dps("<tpEmit>1</tpEmit><prest><CNPJ>12.345.678/0001-99</CNPJ></prest><serv><cServ/></serv>")
    with mock.patch("requests.get", side_effect=get):
        resolvedor = ResolvedorMunicipio(os.path.join(pasta, "dados"), pausa=0)
        assert corrigir_clocemi(pelo_cnpj, resolvedor)[0] == 1 and _tag(pelo_cnpj, "cLocEmi").text == "3106200"
    sem_dados = _dps("<tpEmit>1</tpEmit><prest><xNome>Sem CNPJ</xNome></prest><serv><cServ/></serv>")
    quantidade, mensagens = corrigir_clocprestacao(sem_dados, None)
    assert quantidade == 0 and mensagens and _tag(sem_dados, "locPrest") is None  # não inventa o município


@caso("NFSe Correção de XMLs (lote): salva só as notas alteradas, mantém o xmlns, separa pendentes e grava o histórico")
def _(pasta):
    from modulos.nfse.correcao_xml.engine import (PASTA_CORRIGIDOS, PASTA_PENDENTES, PASTA_PENDENTES_SEM_CEP,
                                                  CorrecaoXmlEngine)

    entrada, saida = os.path.join(pasta, "xmls"), os.path.join(pasta, "saida")
    os.makedirs(entrada)
    corpo = ("<dhEmi>2026-07-20T08:30:00-03:00</dhEmi><dCompet>{compet}</dCompet><tpEmit>1</tpEmit>"
             "<cLocEmi>3106200</cLocEmi><prest><end><endNac><cMun>3106200</cMun><CEP>30130010</CEP></endNac></end></prest>"
             "<toma><xNome>{nome}</xNome><end><endNac>{cmun}<CEP>{cep}</CEP></endNac><xLgr>Rua A</xLgr>"
             "<nro>10</nro></end></toma><serv><locPrest><cLocPrestacao>3106200</cLocPrestacao></locPrest></serv>")
    notas = {
        "corrigir.xml": corpo.format(compet="2026-07-19", nome=_estragar("João"), cmun="", cep="01001000"),
        "certa.xml": corpo.format(compet="2026-07-20", nome="João", cmun="<cMun>3550308</cMun>", cep="01001000"),
        "sem_cep.xml": corpo.format(compet="2026-07-20", nome="João", cmun="", cep="99999999"),
    }
    assinatura = '<Signature xmlns="http://www.w3.org/2000/09/xmldsig#"><SignedInfo/></Signature>'
    for nome, conteudo in notas.items():
        open(os.path.join(entrada, nome), "w", encoding="utf-8").write(
            f'<?xml version="1.0" encoding="UTF-8"?><DPS xmlns="{NS_NFSE}" versao="1.00">'
            f'<infDPS Id="DPS1">{conteudo}</infDPS>{assinatura}</DPS>')
    open(os.path.join(entrada, "quebrada.xml"), "w").write("<DPS><infDPS>")
    open(os.path.join(entrada, "leia.txt"), "w").write("não é xml")

    get, _chamadas = _apis({"viacep.com.br/ws/01001000": Resposta(dados={"ibge": "3550308"}),
                            "viacep.com.br/ws/99999999": Resposta(dados={"erro": True})})
    q = queue.Queue()
    with mock.patch("requests.get", side_effect=get):
        CorrecaoXmlEngine(q, threading.Event(), pasta_entrada=entrada, pasta_saida=saida,
                          log_dir=os.path.join(pasta, "logs"), dados_dir=os.path.join(pasta, "dados"), pausa=0).run()
    evs = eventos(q)
    done = [e for e in evs if e["kind"] in ("done", "fatal_error")][0]
    assert done["kind"] == "done" and done["total"] == 4, done
    assert (done["corrigidos"], done["sem_alteracao"], done["pendentes_cep"], done["pendentes_erro"]) == (1, 1, 1, 1), done
    assert done["por_correcao"]["encoding"] == 1 and done["por_correcao"]["cmun"] == 1
    assert done["por_correcao"]["competencia"] == 1, done

    assert os.listdir(os.path.join(saida, PASTA_CORRIGIDOS)) == ["corrigir.xml"]  # a nota certa não é copiada
    assert os.listdir(os.path.join(saida, PASTA_PENDENTES_SEM_CEP)) == ["sem_cep.xml"]
    assert os.listdir(os.path.join(saida, PASTA_PENDENTES)) == ["quebrada.xml"]
    corrigido = open(os.path.join(saida, PASTA_CORRIGIDOS, "corrigir.xml"), encoding="utf-8").read()
    # namespaces como no original: xmlns na raiz e na assinatura, sem prefixos ns0:/ns1:
    assert corrigido.startswith(f'<?xml version="1.0" encoding="UTF-8"?>\n<DPS xmlns="{NS_NFSE}" versao="1.00">')
    assert '\n  <Signature xmlns="http://www.w3.org/2000/09/xmldsig#">\n    <SignedInfo/>' in corrigido, corrigido[-200:]
    assert "ns0:" not in corrigido and "ns1:" not in corrigido
    assert "<xNome>João</xNome>" in corrigido and "<dCompet>2026-07-20</dCompet>" in corrigido
    # reindentado com 2 espaços, com o cMun criado imediatamente antes do CEP
    assert '\n  <infDPS Id="DPS1">\n    <dhEmi>' in corrigido
    assert "\n          <cMun>3550308</cMun>\n          <CEP>01001000</CEP>\n" in corrigido, corrigido
    import xml.etree.ElementTree as ET
    relido = ET.fromstring(corrigido.encode("utf-8"))  # continua um XML válido, no mesmo namespace
    assert relido.tag == f"{{{NS_NFSE}}}DPS" and _tag(relido, "Signature").tag.startswith("{http://www.w3.org/2000/09")

    # XML com prefixo declarado e caracteres especiais: o prefixo é mantido e o texto é escapado
    from modulos.nfse.correcao_xml.engine import gravar_xml
    com_prefixo = os.path.join(pasta, "prefixo.xml")
    open(com_prefixo, "w", encoding="utf-8").write(
        f'<nfse:DPS xmlns:nfse="{NS_NFSE}"><nfse:xNome a="1 &amp; &quot;2&quot;">A &amp; B &lt; C</nfse:xNome></nfse:DPS>')
    gravar_xml(ET.parse(com_prefixo), com_prefixo, com_prefixo + ".saida")
    regravado = open(com_prefixo + ".saida", encoding="utf-8").read()
    assert f'<nfse:DPS xmlns:nfse="{NS_NFSE}">' in regravado and "A &amp; B &lt; C" in regravado, regravado
    sem_ns = os.path.join(pasta, "sem_ns.xml")
    open(sem_ns, "w", encoding="utf-8").write('<DPS><xNome a="1 &amp; &quot;2&quot;">A &amp; B &lt; C</xNome><v/></DPS>')
    gravar_xml(ET.parse(sem_ns), sem_ns, sem_ns + ".saida")
    regravado = open(sem_ns + ".saida", encoding="utf-8").read()
    assert ET.fromstring(regravado.encode("utf-8")).find("xNome").text == "A & B < C" and "<v/>" in regravado
    assert ET.fromstring(regravado.encode("utf-8")).find("xNome").get("a") == '1 & "2"', regravado
    # os originais continuam como estavam
    assert _estragar("João") in open(os.path.join(entrada, "corrigir.xml"), encoding="utf-8").read()
    historico = open(done["historico"], encoding="utf-8").read()
    for trecho in ("DPS: corrigir.xml", "Status: Corrigida", "Status: Sem alterações",
                   "Pendente: CEP inválido ou sem resolução", "Pendente: XML inválido", "Competência corrigida"):
        assert trecho in historico, trecho

    # só as correções marcadas rodam: sem cMun, a nota sem CEP deixa de ser pendente
    q = queue.Queue()
    with mock.patch("requests.get") as sem_rede:
        CorrecaoXmlEngine(q, threading.Event(), pasta_entrada=entrada, pasta_saida=os.path.join(pasta, "saida2"),
                          log_dir=os.path.join(pasta, "logs"), dados_dir=os.path.join(pasta, "dados2"),
                          correcoes=["encoding", "competencia"], pausa=0).run()
    done = [e for e in eventos(q) if e["kind"] in ("done", "fatal_error")][0]
    assert (done["corrigidos"], done["sem_alteracao"], done["pendentes_cep"]) == (1, 2, 0) and not sem_rede.called, done


def _contraste(cor_a, cor_b):
    from ui.theme import luminancia
    clara, escura = sorted((luminancia(cor_a), luminancia(cor_b)), reverse=True)
    return (clara + 0.05) / (escura + 0.05)


@caso("Animações: a escolha do usuário vale acima do Windows; Área de Trabalho Remota não trava os Pokémon")
def _(pasta):
    from ui import animacao
    casos = [
        # (escolha salva, Windows com animações, sessão remota) -> animações ligadas?
        ((None, True, False), True),
        ((None, False, False), False),   # Windows pediu movimento reduzido
        ((None, False, True), True),     # Área de Trabalho Remota: o Windows só economiza banda
        ((True, False, False), True),    # marcado na janela de temas: vale mesmo com o Windows desligado
        ((False, True, False), False),   # desmarcado: vale sempre
        ((False, False, True), False),
    ]
    for (escolha, windows, remota), esperado in casos:
        with mock.patch.object(animacao.preferencias, "obter", lambda chave, padrao=None: escolha), \
             mock.patch.object(animacao, "movimento_do_sistema_ligado", lambda: windows), \
             mock.patch.object(animacao, "sessao_remota", lambda: remota):
            animacao._estado["ativas"] = None
            assert animacao.ativas() is esperado and animacao.preferencia_ligada() is esperado, \
                (escolha, windows, remota, esperado)
            assert animacao.windows_desliga() is (not windows and not remota)
    animacao._estado["ativas"] = None
    assert isinstance(animacao.sessao_remota(), bool)  # a chamada real ao Windows não falha


@caso("Temas: os prontos têm todas as cores, texto legível e botões com texto de bom contraste")
def _(pasta):
    from ui import temas
    from ui.theme import COLORS, ESTILOS_DE_COR, sobre

    assert temas.PADRAO in temas.PRONTOS and len(temas.PRONTOS) >= 6
    for tema_id, tema in temas.PRONTOS.items():
        cores = tema["cores"]
        assert list(cores) == temas.NOMES_CHAVES, tema_id
        assert all(temas.validar_cor(cor) == cor for cor in cores.values()), tema_id
        paleta = temas.completar(cores)
        assert set(paleta) == set(COLORS), (tema_id, set(paleta) ^ set(COLORS))
        for fundo in ("bg", "panel", "card"):
            assert _contraste(paleta["fg"], paleta[fundo]) >= 7, (tema_id, fundo)
            assert _contraste(paleta["fg_muted"], paleta[fundo]) >= 4, (tema_id, fundo)
            for chave in ("accent", "success", "warning", "danger"):  # texto colorido do log e dos contadores
                assert _contraste(paleta[chave], paleta[fundo]) >= 3, (tema_id, chave, fundo)
        for chave in ESTILOS_DE_COR.values():  # texto dos botões de cor cheia
            assert _contraste(sobre(paleta[chave]), paleta[chave]) >= 4.5, (tema_id, chave)


@caso("Temas: criar, ativar, editar e excluir um tema do usuário; código de cor inválido não entra")
def _(pasta):
    from core import preferencias
    from ui import temas
    from ui.theme import COLORS

    with mock.patch.dict(os.environ, {"MEGAQG_PREFERENCIAS": os.path.join(pasta, "prefs", "preferencias.json")}):
        assert temas.validar_cor("4DA3FF") == "#4da3ff" and temas.validar_cor("#abc") == "#aabbcc"
        assert temas.validar_cor("azul") is None and temas.validar_cor("#12345") is None
        assert temas.ativar("não existe") == temas.PADRAO  # tema desconhecido cai no padrão

        cores = dict(temas.PRONTOS["neon"]["cores"], accent="#ff7a1a", bg="isto não é cor")
        tema_id = temas.salvar_personalizado("  Meu   Laranja ", cores)
        assert tema_id == "meu-meu-laranja" and temas.todos()[tema_id]["nome"] == "Meu Laranja"
        salvo = temas.personalizados()[tema_id]["cores"]
        assert salvo["accent"] == "#ff7a1a" and salvo["bg"] == temas.PRONTOS[temas.PADRAO]["cores"]["bg"]
        assert temas.ativar(tema_id) == tema_id and COLORS["accent"] == "#ff7a1a"
        assert preferencias.obter("tema") == tema_id and COLORS["accent_soft"] != COLORS["accent"]

        assert temas.salvar_personalizado("Meu Laranja", dict(salvo, accent="#00ffaa"), tema_id) == tema_id
        assert temas.salvar_personalizado("Meu Laranja", salvo) == "meu-meu-laranja-2"  # mesmo nome, outro tema
        assert len(temas.personalizados()) == 2
        assert temas.excluir_personalizado(tema_id) and not temas.excluir_personalizado(tema_id)
        assert temas.ativar(tema_id) == temas.PADRAO  # o tema excluído não existe mais

        open(preferencias.caminho(), "w").write("{ arquivo quebrado")
        assert preferencias.carregar() == {} and temas.personalizados() == {}  # não trava o app
        assert temas.ativar(None, salvar=False) == temas.PADRAO
    temas.ativar(temas.PADRAO, salvar=False)


if __name__ == "__main__":
    for rodar in CASOS:
        rodar()
    print(f"\n{len(CASOS) - len(falhas)}/{len(CASOS)} casos passaram.")
    sys.exit(1 if falhas else 0)
